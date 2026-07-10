#!/usr/bin/env python3
"""
Extract SM price lists from Excel files into versioned JSON.
Usage: python extract_price_data.py [version]
  version: optional override, e.g. v26.07.01
  If omitted, auto-detects the newest version from the Price List Ref folder.
"""
import openpyxl, json, os, sys, re, glob
from datetime import date

BASE        = "C:/Projects/Stella Mobili BOM Tool"
PL_DIR      = f"{BASE}/docs/data/Price List Ref"
OUT_DIR     = f"{BASE}/docs/data/price_lists"

if len(sys.argv) > 1:
    VERSION = sys.argv[1]
else:
    matches = glob.glob(f"{PL_DIR}/Kitchen PL SM v*.xlsx")
    if not matches:
        print(f"ERROR: No 'Kitchen PL SM v*.xlsx' found in {PL_DIR}")
        sys.exit(1)
    versions = [m.group(0) for m in (re.search(r'v\d+\.\d+\.\d+', f) for f in matches) if m]
    if not versions:
        print(f"ERROR: Could not parse version from filenames in {PL_DIR}")
        sys.exit(1)
    VERSION = sorted(versions)[-1]
    print(f"Auto-detected version: {VERSION}")
TODAY   = date.today().isoformat()
os.makedirs(OUT_DIR, exist_ok=True)

# ─── helpers ──────────────────────────────────────────────────────────

def fix_enc(s):
    """Fix mojibake: UTF-8 bytes stored as Latin-1 chars in Excel cells."""
    if not s:
        return s
    try:
        return s.encode('latin-1').decode('utf-8')
    except (UnicodeEncodeError, UnicodeDecodeError):
        return s

def sheet_rows(wb, sheet):
    ws = wb[sheet]
    it = ws.iter_rows(values_only=True)
    hdrs = [fix_enc(str(h).strip()) if h is not None else '' for h in next(it)]
    out = []
    for row in it:
        d = {}
        for i, v in enumerate(row):
            if i < len(hdrs) and hdrs[i]:
                d[hdrs[i]] = fix_enc(str(v).strip()) if v is not None else ''
        if any(v for v in d.values()):
            out.append(d)
    return out

def sv(r, k, default=''):
    v = r.get(k, default)
    s = str(v).strip() if v is not None else ''
    return s if s and s != 'None' else default

def iv(r, k):
    v = r.get(k, '')
    try:
        f = float(str(v).strip())
        return int(f) if str(v).strip() not in ('', 'None') else None
    except Exception:
        return None

def pv(r, k='PRICE'):
    try:
        return int(float(str(r.get(k, 0))))
    except Exception:
        return 0

def is_active(r):
    return sv(r, 'ACTIVE', 'YES').upper() not in ('NO', '0', 'TIDAK', 'FALSE')

# ─── KITCHEN ──────────────────────────────────────────────────────────

def extract_kitchen(version):
    path = f"{PL_DIR}/Kitchen PL SM {version}.xlsx"
    wb   = openpyxl.load_workbook(path, read_only=True)
    out  = {}

    # Standard sheets: Base Unit / Wall Unit / Tall Unit
    for sheet_name, top_cat in [
        ('BASE UNIT (FLAT)',  'Base Unit'),
        ('WALL UNIT (FLAT)',  'Wall Unit'),
        ('TALL UNIT (FLAT)',  'Tall Unit'),
    ]:
        groups = {}
        for r in sheet_rows(wb, sheet_name):
            if not is_active(r): continue
            desc = sv(r, 'DESCRIPTION')
            if not desc: continue
            pr = pv(r)
            if not pr: continue
            subcat = sv(r, 'CATEGORY')
            w = iv(r, 'WIDTH'); h = iv(r, 'HEIGHT'); d = iv(r, 'DEPTH')
            finish = sv(r, 'FINISH') or '(No Finish)'
            handle = sv(r, 'HANDLE', 'Standard').lower() or 'standard'
            thick  = sv(r, 'THICKNESS')
            unit   = sv(r, 'UNIT', 'set')
            notes  = sv(r, 'NOTES')
            gkey = (subcat, desc, w, h, d)
            if gkey not in groups:
                groups[gkey] = {'subcat': subcat, 'desc': desc,
                                 'w': w, 'h': h, 'd': d,
                                 'thickness': thick, 'unit': unit,
                                 'notes': notes, 'prices': {}}
            groups[gkey]['prices'][f"{finish}__{handle}"] = pr
            if thick and not groups[gkey]['thickness']:
                groups[gkey]['thickness'] = thick
        out[top_cat] = list(groups.values())
        print(f"  Kitchen/{top_cat}: {len(out[top_cat])} items")

    # Artdeco (FLAT) — SECTION col = BASE UNIT / WALL UNIT / TALL UNIT
    sec_map = {'BASE UNIT': 'Artdeco Base Unit',
               'WALL UNIT': 'Artdeco Wall Unit',
               'TALL UNIT': 'Artdeco Tall Unit'}
    ad_groups = {}
    for r in sheet_rows(wb, 'ARTDECO (FLAT)'):
        if not is_active(r): continue
        desc = sv(r, 'DESCRIPTION')
        if not desc: continue
        pr = pv(r)
        if not pr: continue
        top_cat = sec_map.get(sv(r, 'SECTION').upper(), 'Artdeco Base Unit')
        subcat  = sv(r, 'CATEGORY')
        w = iv(r, 'WIDTH'); h = iv(r, 'HEIGHT'); d = iv(r, 'DEPTH')
        finish  = sv(r, 'FINISH') or '(No Finish)'
        handle  = sv(r, 'HANDLE', 'Standard').lower() or 'standard'
        thick   = sv(r, 'THICKNESS')
        unit    = sv(r, 'UNIT', 'set')
        notes   = sv(r, 'NOTES')
        gkey = (top_cat, subcat, desc, w, h, d)
        if gkey not in ad_groups:
            ad_groups[gkey] = {'subcat': subcat, 'desc': desc,
                                'w': w, 'h': h, 'd': d,
                                'thickness': thick, 'unit': unit,
                                'notes': notes, 'prices': {}}
        ad_groups[gkey]['prices'][f"{finish}__{handle}"] = pr
        if thick and not ad_groups[gkey]['thickness']:
            ad_groups[gkey]['thickness'] = thick
    for tc in ['Artdeco Base Unit', 'Artdeco Wall Unit', 'Artdeco Tall Unit']:
        out[tc] = [v for (t, *_), v in ad_groups.items() if t == tc]
        print(f"  Kitchen/{tc}: {len(out[tc])} items")

    # Glass Doors (FLAT)
    gd_groups = {}
    for r in sheet_rows(wb, 'WU TU Glass Doors (FLAT)'):
        if not is_active(r): continue
        desc = sv(r, 'DESCRIPTION')
        if not desc: continue
        pr = pv(r)
        if not pr: continue
        subcat = sv(r, 'CATEGORY')
        w = iv(r, 'WIDTH'); h = iv(r, 'HEIGHT'); d = iv(r, 'DEPTH')
        glass = sv(r, 'GLASS TYPE'); hinge = sv(r, 'HINGE TYPE')
        unit  = sv(r, 'UNIT', 'set'); notes = sv(r, 'NOTES')
        gkey = (subcat, desc, w, h, d)
        if gkey not in gd_groups:
            gd_groups[gkey] = {'subcat': subcat, 'desc': desc,
                                'w': w, 'h': h, 'd': d,
                                'thickness': '', 'unit': unit,
                                'notes': notes, 'prices': {}}
        gd_groups[gkey]['prices'][f"{glass} | {hinge}__standard"] = pr
    out['Glass Doors'] = list(gd_groups.values())
    print(f"  Kitchen/Glass Doors: {len(out['Glass Doors'])} items")

    # Ornament Artdeco (FLAT)
    oa_groups = {}
    for r in sheet_rows(wb, 'Ornament Artdeco (FLAT)'):
        if not is_active(r): continue
        desc = sv(r, 'DESCRIPTION')
        if not desc: continue
        pr = pv(r)
        if not pr: continue
        subcat = sv(r, 'CATEGORY')
        w = iv(r, 'WIDTH'); h = iv(r, 'HEIGHT'); d = iv(r, 'DEPTH')
        color = sv(r, 'COLOR') or '(No Finish)'
        unit  = sv(r, 'UNIT', 'set'); notes = sv(r, 'NOTES')
        gkey = (subcat, desc, w, h, d)
        if gkey not in oa_groups:
            oa_groups[gkey] = {'subcat': subcat, 'desc': desc,
                                'w': w, 'h': h, 'd': d,
                                'thickness': '', 'unit': unit,
                                'notes': notes, 'prices': {}}
        oa_groups[gkey]['prices'][f"{color}__standard"] = pr
    out['Ornament Artdeco'] = list(oa_groups.values())
    print(f"  Kitchen/Ornament Artdeco: {len(out['Ornament Artdeco'])} items")

    wb.close()
    return out

# ─── WARDROBE ─────────────────────────────────────────────────────────

def extract_wardrobe(version):
    path = f"{PL_DIR}/Wardrobe PL SM {version}.xlsx"
    wb   = openpyxl.load_workbook(path, read_only=True)
    out  = {}

    # Finish-based panels — SECTION col gives category name
    for sheet_name in ['SWINGDOOR110 (FLAT)', 'SWINGDOOR170 (FLAT)', 'SLIDING (FLAT)']:
        groups = {}; section_name = None
        for r in sheet_rows(wb, sheet_name):
            if not is_active(r): continue
            desc = sv(r, 'DESCRIPTION')
            if not desc: continue
            pr = pv(r)
            if not pr: continue
            if not section_name: section_name = sv(r, 'SECTION')
            subcat = sv(r, 'CATEGORY')
            finish = sv(r, 'FINISH') or '(No Finish)'
            thick  = sv(r, 'THICKNESS')
            unit   = sv(r, 'UNIT', 'set'); notes = sv(r, 'NOTES')
            gkey = (subcat, desc)
            if gkey not in groups:
                groups[gkey] = {'subcat': subcat, 'desc': desc,
                                 'unit': unit, 'notes': notes,
                                 'prices': {}, 'thicknesses': {}}
            pk = f"{finish}__standard"
            groups[gkey]['prices'][pk] = pr
            if thick: groups[gkey]['thicknesses'][pk] = thick
            if notes and not groups[gkey]['notes']:
                groups[gkey]['notes'] = notes
        if section_name:
            items = list(groups.values())
            for it in items:
                if not it['thicknesses']: del it['thicknesses']
            out[section_name] = items
            print(f"  Wardrobe/{section_name}: {len(items)} items")

    # Swing Glass Door
    sg_groups = {}; sg_section = None
    for r in sheet_rows(wb, 'SWINGGLASS (FLAT)'):
        if not is_active(r): continue
        desc = sv(r, 'DESCRIPTION')
        if not desc: continue
        pr = pv(r)
        if not pr: continue
        if not sg_section: sg_section = sv(r, 'SECTION')
        subcat = sv(r, 'CATEGORY')
        hinge = sv(r, 'HINGE TYPE'); glass = sv(r, 'GLASS TYPE')
        unit  = sv(r, 'UNIT', 'set'); notes = sv(r, 'NOTES')
        gkey = (subcat, desc)
        if gkey not in sg_groups:
            sg_groups[gkey] = {'subcat': subcat, 'desc': desc,
                                'unit': unit, 'notes': notes, 'prices': {}}
        sg_groups[gkey]['prices'][f"{glass} | {hinge}__standard"] = pr
    if sg_section:
        out[sg_section] = list(sg_groups.values())
        print(f"  Wardrobe/{sg_section}: {len(out[sg_section])} items")

    # Sliding Glass Door
    slg_groups = {}; slg_section = None
    for r in sheet_rows(wb, 'SLIDINGGLASS (FLAT)'):
        if not is_active(r): continue
        desc = sv(r, 'DESCRIPTION')
        if not desc: continue
        pr = pv(r)
        if not pr: continue
        if not slg_section: slg_section = sv(r, 'SECTION')
        subcat = sv(r, 'CATEGORY')
        frame = sv(r, 'FRAME TYPE'); glass = sv(r, 'GLASS TYPE')
        unit  = sv(r, 'UNIT', 'set'); notes = sv(r, 'NOTES')
        gkey = (subcat, desc)
        if gkey not in slg_groups:
            slg_groups[gkey] = {'subcat': subcat, 'desc': desc,
                                  'unit': unit, 'notes': notes, 'prices': {}}
        slg_groups[gkey]['prices'][f"{glass} | {frame}__standard"] = pr
    if slg_section:
        out[slg_section] = list(slg_groups.values())
        print(f"  Wardrobe/{slg_section}: {len(out[slg_section])} items")

    # Upper Cabinet (FINISH + DEPTH — separate items per depth)
    uc_groups = {}; uc_section = None
    for r in sheet_rows(wb, 'UPPERCAB (FLAT)'):
        if not is_active(r): continue
        desc = sv(r, 'DESCRIPTION')
        if not desc: continue
        pr = pv(r)
        if not pr: continue
        if not uc_section: uc_section = sv(r, 'SECTION')
        subcat = sv(r, 'CATEGORY')
        finish = sv(r, 'FINISH') or '(No Finish)'
        depth  = sv(r, 'DEPTH')
        unit   = sv(r, 'UNIT', 'set'); notes = sv(r, 'NOTES')
        gkey = (subcat, desc, depth)
        if gkey not in uc_groups:
            uc_groups[gkey] = {'subcat': subcat, 'desc': desc, 'depth': depth,
                                'unit': unit, 'notes': notes, 'prices': {}}
        uc_groups[gkey]['prices'][f"{finish}__standard"] = pr
    if uc_section:
        out[uc_section] = list(uc_groups.values())
        print(f"  Wardrobe/{uc_section}: {len(out[uc_section])} items")

    # Upper Cabinet Mirror (GLASS TYPE + HINGE TYPE + DEPTH)
    ucm_groups = {}; ucm_section = None
    for r in sheet_rows(wb, 'UPPERCABMIRROR (FLAT)'):
        if not is_active(r): continue
        desc = sv(r, 'DESCRIPTION')
        if not desc: continue
        pr = pv(r)
        if not pr: continue
        if not ucm_section: ucm_section = sv(r, 'SECTION')
        subcat = sv(r, 'CATEGORY')
        hinge = sv(r, 'HINGE TYPE'); glass = sv(r, 'GLASS TYPE')
        depth  = sv(r, 'DEPTH')
        unit   = sv(r, 'UNIT', 'set'); notes = sv(r, 'NOTES')
        gkey = (subcat, desc, depth)
        if gkey not in ucm_groups:
            ucm_groups[gkey] = {'subcat': subcat, 'desc': desc, 'depth': depth,
                                  'unit': unit, 'notes': notes, 'prices': {}}
        ucm_groups[gkey]['prices'][f"{glass} | {hinge}__standard"] = pr
    if ucm_section:
        out[ucm_section] = list(ucm_groups.values())
        print(f"  Wardrobe/{ucm_section}: {len(out[ucm_section])} items")

    # Base Cabinet BLUM + STD (W/H/D + FINISH)
    for sheet_name, section_name in [
        ('BASEINTERNAL BLUM (FLAT)', 'Base Cabinet BLUM'),
        ('BASEINTERNAL STD (FLAT)',  'Base Cabinet STD'),
    ]:
        groups = {}
        for r in sheet_rows(wb, sheet_name):
            if not is_active(r): continue
            desc = sv(r, 'DESCRIPTION')
            if not desc: continue
            pr = pv(r)
            if not pr: continue
            subcat = sv(r, 'CATEGORY')
            w = iv(r, 'WIDTH'); h = iv(r, 'HEIGHT'); d = iv(r, 'DEPTH')
            finish = sv(r, 'FINISH') or '(No Finish)'
            unit   = sv(r, 'UNIT', 'set'); notes = sv(r, 'NOTES')
            gkey = (subcat, desc, w, h, d)
            if gkey not in groups:
                groups[gkey] = {'subcat': subcat, 'desc': desc,
                                 'w': w, 'h': h, 'd': d,
                                 'unit': unit, 'notes': notes, 'prices': {}}
            groups[gkey]['prices'][f"{finish}__standard"] = pr
        out[section_name] = list(groups.values())
        print(f"  Wardrobe/{section_name}: {len(out[section_name])} items")

    # Marble Door Frame (W/H, no finish)
    mdf_groups = {}; mdf_section = None
    for r in sheet_rows(wb, 'MARBLEDOORFRAME (FLAT)'):
        if not is_active(r): continue
        desc = sv(r, 'DESCRIPTION')
        if not desc: continue
        pr = pv(r)
        if not pr: continue
        if not mdf_section: mdf_section = sv(r, 'SECTION')
        subcat = sv(r, 'CATEGORY')
        w = iv(r, 'WIDTH'); h = iv(r, 'HEIGHT')
        thick  = sv(r, 'THICKNESS')
        unit   = sv(r, 'UNIT', 'set'); notes = sv(r, 'NOTES')
        gkey = (subcat, desc, w, h)
        if gkey not in mdf_groups:
            mdf_groups[gkey] = {'subcat': subcat, 'desc': desc,
                                  'w': w, 'h': h, 'thickness': thick,
                                  'unit': unit, 'notes': notes,
                                  'prices': {'(No Finish)__standard': pr}}
    if mdf_section:
        out[mdf_section] = list(mdf_groups.values())
        print(f"  Wardrobe/{mdf_section}: {len(out[mdf_section])} items")

    # Wardrobe Component (W/H/D, no finish)
    wdc_groups = {}; wdc_section = None
    for r in sheet_rows(wb, 'WD COMP (FLAT)'):
        if not is_active(r): continue
        desc = sv(r, 'DESCRIPTION')
        if not desc: continue
        pr = pv(r)
        if not pr: continue
        if not wdc_section: wdc_section = sv(r, 'SECTION')
        subcat = sv(r, 'CATEGORY')
        w = iv(r, 'WIDTH'); h = iv(r, 'HEIGHT'); d = iv(r, 'DEPTH')
        thick  = sv(r, 'THICKNESS')
        unit   = sv(r, 'UNIT', 'pcs'); notes = sv(r, 'NOTES')
        gkey = (subcat, desc, w, h, d)
        if gkey not in wdc_groups:
            wdc_groups[gkey] = {'subcat': subcat, 'desc': desc,
                                  'w': w, 'h': h, 'd': d,
                                  'thickness': thick if thick else '',
                                  'unit': unit, 'notes': notes,
                                  'prices': {'(No Finish)__standard': pr}}
    if wdc_section:
        out[wdc_section] = list(wdc_groups.values())
        print(f"  Wardrobe/{wdc_section}: {len(out[wdc_section])} items")

    wb.close()
    return out

# ─── ACCESSORIES ──────────────────────────────────────────────────────

def extract_accessories(version):
    path = f"{PL_DIR}/Accessories PL SM {version}.xlsx"
    wb   = openpyxl.load_workbook(path, read_only=True)
    out  = {}

    # Solid Surface (THICKNESS mm + DEPTH mm as dimensions)
    ss_groups = {}
    for r in sheet_rows(wb, 'Solid Surface (FLAT)'):
        if not is_active(r): continue
        desc = sv(r, 'DESCRIPTION')
        if not desc: continue
        pr = pv(r)
        if not pr: continue
        subcat   = sv(r, 'CATEGORY')
        thick_mm = iv(r, 'THICKNESS (mm)')
        depth_mm = iv(r, 'DEPTH (mm)')
        unit     = sv(r, 'UNIT', 'meter'); notes = sv(r, 'NOTES')
        gkey = (subcat, desc, thick_mm, depth_mm)
        if gkey not in ss_groups:
            ss_groups[gkey] = {'subcat': subcat, 'desc': desc,
                                'thickness_mm': thick_mm, 'depth_mm': depth_mm,
                                'unit': unit, 'notes': notes, 'prices': {}}
        ss_groups[gkey]['prices']['(No Finish)__standard'] = pr
    out['Solid Surface'] = list(ss_groups.values())
    print(f"  Accessories/Solid Surface: {len(out['Solid Surface'])} items")

    # Simple flat sheets
    def flat_sheet(sheet_name, top_cat, default_unit='set', price_col='PRICE'):
        groups = {}
        for r in sheet_rows(wb, sheet_name):
            if not is_active(r): continue
            desc = sv(r, 'DESCRIPTION')
            if not desc: continue
            pr = pv(r, price_col)
            if not pr: continue
            subcat = sv(r, 'CATEGORY')
            unit   = sv(r, 'UNIT', default_unit) if 'UNIT' in r else default_unit
            notes  = sv(r, 'NOTES')
            gkey = (subcat, desc)
            if gkey not in groups:
                groups[gkey] = {'subcat': subcat, 'desc': desc,
                                 'unit': unit, 'notes': notes, 'prices': {}}
            groups[gkey]['prices']['(No Finish)__standard'] = pr
        out[top_cat] = list(groups.values())
        print(f"  Accessories/{top_cat}: {len(out[top_cat])} items")

    flat_sheet('Lighting (FLAT)',     'Lighting',             'meter')
    flat_sheet('Wardrobe Acc (FLAT)', 'Wardrobe Accessories', 'set')
    flat_sheet('Kitchen Acc (FLAT)',  'Kitchen Accessories',  'set')
    flat_sheet('Mirror LED (FLAT)',   'Mirror',               'm2')

    # Handles (CODE + LENGTH as metadata)
    h_groups = {}
    for r in sheet_rows(wb, 'Handles (FLAT)'):
        if not is_active(r): continue
        desc = sv(r, 'DESCRIPTION')
        if not desc: continue
        pr = pv(r, 'PRICE')
        if not pr: continue
        subcat = sv(r, 'CATEGORY')
        code   = sv(r, 'CODE'); length = sv(r, 'LENGTH')
        notes  = sv(r, 'NOTES')
        # Append length to description so dropdown shows e.g. "Handle Kuku Macan 180mm"
        length_mm = f"{length}mm" if length else ''
        if length_mm and length_mm not in desc:
            desc_full = f"{desc} {length_mm}"
        else:
            desc_full = desc
        gkey = (subcat, desc_full, code, length)
        if gkey not in h_groups:
            h_groups[gkey] = {'subcat': subcat, 'desc': desc_full,
                               'code': code, 'length': length,
                               'unit': 'pcs', 'notes': notes, 'prices': {}}
        h_groups[gkey]['prices']['(No Finish)__standard'] = pr
    out['Handles'] = list(h_groups.values())
    print(f"  Accessories/Handles: {len(out['Handles'])} items")

    wb.close()
    return out

# ─── MAIN ─────────────────────────────────────────────────────────────

print(f"\nExtracting price list: {VERSION}\n")

print("[Kitchen]")
kitchen_out = {'version': VERSION, 'type': 'kitchen',
               'extractedAt': TODAY, 'items': extract_kitchen(VERSION)}

print("\n[Wardrobe]")
wardrobe_out = {'version': VERSION, 'type': 'wardrobe',
                'extractedAt': TODAY, 'items': extract_wardrobe(VERSION)}

print("\n[Accessories]")
acc_out = {'version': VERSION, 'type': 'accessories',
           'extractedAt': TODAY, 'items': extract_accessories(VERSION)}

print()
for name, data in [('kitchen', kitchen_out), ('wardrobe', wardrobe_out),
                   ('accessories', acc_out)]:
    fpath = f"{OUT_DIR}/{name}_{VERSION}.json"
    with open(fpath, 'w', encoding='utf-8') as f:
        json.dump(data, f, ensure_ascii=False, separators=(',', ':'))
    print(f"Written: {fpath}  ({os.path.getsize(fpath)/1024:.0f} KB)")

# Update index.json
idx_path = f"{OUT_DIR}/index.json"
if os.path.exists(idx_path):
    with open(idx_path, encoding='utf-8') as f:
        idx = json.load(f)
else:
    idx = {'latest': VERSION, 'versions': []}
if VERSION not in idx['versions']:
    idx['versions'].append(VERSION)
    idx['versions'].sort(reverse=True)
idx['latest'] = idx['versions'][0]
with open(idx_path, 'w', encoding='utf-8') as f:
    json.dump(idx, f, ensure_ascii=False, indent=2)
print(f"Updated: {idx_path}  versions={idx['versions']}")
print("\nDone.")
